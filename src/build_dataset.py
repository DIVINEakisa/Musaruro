"""Build tidy Musaruro datasets from NISR Seasonal Agricultural Survey (SAS) 2026 tables.

Inputs  : data/raw/Tables_2026_Season_A.xlsx, Tables_2026_Season_B.xlsx
Outputs : data/processed/*.csv  (+ data_quality_log.csv)

Every number in the outputs is copied from an NISR table; nothing is imputed.
Blank cells in the source stay blank (NaN). Run:  python src/build_dataset.py
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed"

PROVINCE = {
    "Nyarugenge": "Kigali City", "Gasabo": "Kigali City", "Kicukiro": "Kigali City",
    "Nyanza": "Southern", "Gisagara": "Southern", "Nyaruguru": "Southern", "Huye": "Southern",
    "Nyamagabe": "Southern", "Ruhango": "Southern", "Muhanga": "Southern", "Kamonyi": "Southern",
    "Karongi": "Western", "Rutsiro": "Western", "Rubavu": "Western", "Nyabihu": "Western",
    "Ngororero": "Western", "Rusizi": "Western", "Nyamasheke": "Western",
    "Rulindo": "Northern", "Gakenke": "Northern", "Musanze": "Northern", "Burera": "Northern",
    "Gicumbi": "Northern",
    "Rwamagana": "Eastern", "Nyagatare": "Eastern", "Gatsibo": "Eastern", "Kayonza": "Eastern",
    "Kirehe": "Eastern", "Ngoma": "Eastern", "Bugesera": "Eastern",
}
DISTRICTS = list(PROVINCE)

# Table numbers differ between the two workbooks.
SEASONS = {
    "A": dict(file="Tables_2026_Season_A.xlsx", area="Table 11", harv="Table 12", yld="Table 13",
              prod="Table 15", seed="Table 20", prac="Table 10", land="Table 9"),
    "B": dict(file="Tables_2026_Season_B.xlsx", area="Table 9", harv="Table 10", yld="Table 11",
              prod="Table 13", seed="Table 18", prac="Table 8", land="Table 7"),
}

# Aggregate columns in the NISR tables (not individual crops).
GROUPS = {"Cereals", "Tubers and Roots", "Banana", "Legumes and Pulses", "Vegetables and Fruits",
          "All crops (total developed land)"}
CROP_ALIASES = {
    "Yams & Taro": "Yams & Taro", "Groundnuts": "Groundnut", "Bananas": "Banana",
    "Bean": "Bean (all)", "Vegetables": "Vegetables", "Fruits": "Fruits",
}
NATIONAL_LABELS = {"national", "total", "s/total"}


def norm(s) -> str:
    s = str(s).replace("&amp;", "&")
    return re.sub(r"\s+", " ", s).strip()


CANON = {n.lower(): n for n in [
    "Cereals", "Maize", "Sorghum", "Paddy rice", "Wheat", "Other cereals", "Tubers and Roots", "Cassava",
    "Sweet potato", "Irish potato", "Yams & Taro", "Banana", "Cooking banana", "Dessert banana",
    "Banana for beer", "Legumes and Pulses", "Bean (all)", "Bush bean", "Climbing bean", "Pea",
    "Groundnut", "Soybean", "Vegetables and Fruits", "Vegetables", "Fruits", "Fodder crops", "Other crops",
    "All crops (total developed land)"]}


def crop_name(raw: str) -> str:
    n = norm(raw)
    n = CROP_ALIASES.get(n, n)
    if n.lower() == "total developed land":
        n = "All crops (total developed land)"
    return CANON.get(n.lower(), n)


def parse_crop_table(ws) -> dict[str, dict[str, float]]:
    """Return {row label: {crop: value}} for a district x crop NISR table."""
    header = None
    for row in ws.iter_rows():
        if any(c.value is not None and norm(c.value) == "Cereals" for c in row):
            header = {c.column: crop_name(c.value) for c in row if c.value is not None and c.column > 2}
            hdr_row = row[0].row
            break
    if header is None:
        raise ValueError(f"No crop header in {ws.title}")
    out: dict[str, dict[str, float]] = {}
    for row in ws.iter_rows(min_row=hdr_row + 1):
        label = row[1].value if len(row) > 1 else None
        if label is None:
            continue
        label = norm(label)
        if label.lower().startswith("source"):
            break
        vals = {}
        for c in row:
            if c.column in header and isinstance(c.value, (int, float)):
                vals[header[c.column]] = float(c.value)
        out[label] = vals
    return out


def rows_by_label(ws, first_col: int, last_col: int) -> dict[str, list]:
    out = {}
    for row in ws.iter_rows():
        label = row[1].value if len(row) > 1 else None
        if label is None:
            continue
        label = norm(label)
        if label.lower().startswith("source"):
            break
        vals = [row[i - 1].value if i - 1 < len(row) else None for i in range(first_col, last_col + 1)]
        out[label] = [float(v) if isinstance(v, (int, float)) else None for v in vals]
    return out


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    crop_rows, nat_rows, seed_rows, prac_rows, land_rows, dq = [], [], [], [], [], []

    for season, cfg in SEASONS.items():
        wb = openpyxl.load_workbook(RAW / cfg["file"], data_only=True)
        tabs = {k: parse_crop_table(wb[cfg[k]]) for k in ("area", "harv", "yld", "prod")}
        names = {"area": "cultivated_ha", "harv": "harvested_ha", "yld": "yield_kg_ha", "prod": "production_mt"}

        # --- crop x district x season -------------------------------------------------
        for district in DISTRICTS:
            for key, tab in tabs.items():
                if district not in tab:
                    dq.append(dict(season=season, table=cfg[key], issue=f"district {district} missing"))
            crops = set().union(*[set(tabs[k].get(district, {})) for k in tabs])
            for crop in crops:
                crop_rows.append(dict(
                    season=season, district=district, province=PROVINCE[district], crop=crop,
                    is_group=crop in GROUPS,
                    **{names[k]: tabs[k].get(district, {}).get(crop) for k in tabs}))

        # --- national reference rows --------------------------------------------------
        for key, tab in tabs.items():
            for label, vals in tab.items():
                if label.lower() in NATIONAL_LABELS:
                    for crop, v in vals.items():
                        nat_rows.append(dict(season=season, crop=crop, measure=names[key], value=v,
                                             source_table=f"{cfg['file']} {cfg[key]}"))

        # --- improved seed by district --------------------------------------------------
        seed = rows_by_label(wb[cfg["seed"]], 3, 11)
        for label, v in seed.items():
            if label in PROVINCE or label.lower() in NATIONAL_LABELS:
                seed_rows.append(dict(
                    season=season, district=label if label in PROVINCE else "National",
                    farmers_improved_pct=v[0], farmers_improved_pct_ssf=v[1], farmers_improved_pct_lsf=v[2],
                    plots_improved_pct=v[3], plots_improved_pct_ssf=v[4], plots_improved_pct_lsf=v[5],
                    land_improved_pct=v[6], land_improved_pct_ssf=v[7], land_improved_pct_lsf=v[8]))

        # --- agricultural practices (ha) ------------------------------------------------
        prac = rows_by_label(wb[cfg["prac"]], 3, 7)
        for label, v in prac.items():
            if label in PROVINCE or label.lower() in NATIONAL_LABELS:
                prac_rows.append(dict(
                    season=season, district=label if label in PROVINCE else "National",
                    irrigated_ha=v[0], erosion_control_ha=v[1], agroforestry_ha=v[2],
                    inorganic_fert_ha=v[3], organic_fert_ha=v[4]))

        # --- land use (000 ha) ----------------------------------------------------------
        land = rows_by_label(wb[cfg["land"]], 3, 12)
        for label, v in land.items():
            if label in PROVINCE or label.lower() in NATIONAL_LABELS:
                land_rows.append(dict(
                    season=season, district=label if label in PROVINCE else "National",
                    total_land_kha=v[0], agricultural_land_kha=v[1], pct_agricultural=v[2],
                    arable_kha=v[3], physical_cultivated_kha=v[4], seasonal_crops_kha=v[5],
                    permanent_crops_kha=v[6], temp_fallow_kha=v[7]))

    crops = pd.DataFrame(crop_rows).sort_values(["season", "district", "crop"]).reset_index(drop=True)
    nat = pd.DataFrame(nat_rows)
    nat = nat.pivot_table(index=["season", "crop"], columns="measure", values="value").reset_index()
    nat.columns.name = None

    # Source-data quality flag: Season A "plots overall" column duplicates "land overall".
    seed = pd.DataFrame(seed_rows)
    a = seed[(seed.season == "A") & (seed.district != "National")]
    dup = (a.plots_improved_pct - a.land_improved_pct).abs().lt(1e-6).mean()
    if dup > 0.9:
        dq.append(dict(season="A", table="Table 20",
                       issue=f"'plots overall' equals 'land overall' in {dup:.0%} of districts and exceeds both SSF "
                             "and LSF values in some (impossible for a weighted average). plots_improved_pct is set "
                             "to NaN for Season A; use farmers_improved_pct as the primary adoption metric."))
        seed.loc[seed.season == "A", "plots_improved_pct"] = float("nan")

    # Consistency check: district values should add up to NISR's national row.
    for measure in ("cultivated_ha", "harvested_ha", "production_mt"):
        s = crops.groupby(["season", "crop"])[measure].sum()
        nref = nat.set_index(["season", "crop"])[measure]
        rel = ((s - nref).abs() / nref).dropna()
        for (season, crop), r in rel[(rel > 0.005) & (nref.reindex(rel.index) > 100)].items():
            dq.append(dict(season=season, table="district vs national",
                           issue=f"{crop} {measure}: districts sum to {s[(season, crop)]:,.0f} but NISR national row "
                                 f"is {nref[(season, crop)]:,.0f} ({r:.1%} gap). Source tables disagree; use district values."))

    # --- survey segments per district (Table 4, same design in both seasons) -------------
    ws4 = openpyxl.load_workbook(RAW / SEASONS["A"]["file"], data_only=True)["Table 4"]
    seg = [dict(district=norm(r[1]), survey_segments=int(r[6])) for r in ws4.iter_rows(values_only=True)
           if r[1] and norm(r[1]) in PROVINCE and isinstance(r[6], (int, float))]
    pd.DataFrame(seg).to_csv(OUT / "survey_segments.csv", index=False)

    crops.to_csv(OUT / "crop_district_season.csv", index=False)
    nat.to_csv(OUT / "national_reference.csv", index=False)
    seed.to_csv(OUT / "improved_seed_district.csv", index=False)
    pd.DataFrame(prac_rows).to_csv(OUT / "practices_district.csv", index=False)
    pd.DataFrame(land_rows).to_csv(OUT / "land_use_district.csv", index=False)
    pd.DataFrame(dq, columns=["season", "table", "issue"]).to_csv(OUT / "data_quality_log.csv", index=False)
    print(f"crop rows {len(crops)} | national rows {len(nat)} | seed rows {len(seed)} | flags {len(dq)}")


if __name__ == "__main__":
    build()

"""Water and soil tables: irrigation by district (SAS 2026 A and B) and province context (AHS 2024).

Run after build_dataset.py:  python -m src.build_water
Outputs (data/processed/): irrigation_district.csv, ahs_water_province.csv, ahs_soil_province.csv
Note: technique and water-source shares are % of IRRIGATED PLOTS, a small sample in most districts. The app says so.
"""
from __future__ import annotations

import difflib

import openpyxl
import pandas as pd

from src.build_dataset import DISTRICTS, OUT, RAW, norm

TABLES = {"A": dict(farm="Table 31", tech="Table 32", src="Table 33"),
          "B": dict(farm="Table 29", tech="Table 30", src="Table 31")}
FILE = {"A": "Tables_2026_Season_A.xlsx", "B": "Tables_2026_Season_B.xlsx"}
TECH = ["surface", "flood", "drip", "sprinkler", "pivot", "traditional"]
SRC = ["rainwater", "treatment_plant", "underground", "lake_stream", "catchment", "other_source"]
PROVINCES = ["Kigali", "South", "West", "North", "East", "Rwanda"]


def match_district(label: str):
    label = norm(label)
    if label in DISTRICTS:
        return label
    m = difflib.get_close_matches(label, DISTRICTS, n=1, cutoff=0.85)
    return m[0] if m else None


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def district_rows(ws, ncols: int, first_col: int = 3):
    out = {}
    for row in ws.iter_rows():
        if len(row) < 2 or row[1].value is None:
            continue
        d = match_district(str(row[1].value))
        if d:
            out[d] = [num(row[first_col - 1 + i].value) if first_col - 1 + i < len(row) else None for i in range(ncols)]
    return out


def province_rows(ws, ncols: int, first_col: int = 3):
    out = {}
    for row in ws.iter_rows():
        if len(row) < 2 or row[1].value is None:
            continue
        label = str(row[1].value).strip()
        if label.startswith("Source"):
            break
        out[label] = [num(row[first_col - 1 + i].value) if first_col - 1 + i < len(row) else None for i in range(ncols)]
    return out


def build() -> None:
    rows = []
    for season, tb in TABLES.items():
        wb = openpyxl.load_workbook(RAW / FILE[season], data_only=True)
        farm = district_rows(wb[tb["farm"]], 9)          # erosion x3, mechanical x3, irrigation x3
        tech = district_rows(wb[tb["tech"]], 6)
        src = district_rows(wb[tb["src"]], 6)
        for d in DISTRICTS:
            f = farm.get(d, [None] * 9)
            r = dict(season=season, district=d, farmers_irrigating_pct=f[6], farmers_irrigating_pct_ssf=f[7],
                     farmers_irrigating_pct_lsf=f[8])
            r.update({"tech_" + k: v for k, v in zip(TECH, tech.get(d, [None] * 6))})
            r.update({"src_" + k: v for k, v in zip(SRC, src.get(d, [None] * 6))})
            rows.append(r)
    irr = pd.DataFrame(rows)
    irr.to_csv(OUT / "irrigation_district.csv", index=False)

    ahs = openpyxl.load_workbook(RAW / "AHS_2024_Tables.xlsx", data_only=True)
    w25 = province_rows(ahs["Table 25"], 7)
    w26 = province_rows(ahs["Table 26"], 4)
    recs = []
    for p in PROVINCES:
        k25 = next((k for k in w25 if k.strip() == p), None)
        k26 = next((k for k in w26 if k.strip() == p), None)
        v25, v26 = w25.get(k25, [None] * 7), w26.get(k26, [None] * 4)
        recs.append(dict(province=p, src_rainwater=v25[0], src_treatment_plant=v25[1], src_underground=v25[2],
                         src_lake_stream=v25[3], src_dams=v25[4], src_other=v25[5], irrigating_households_000=v25[6],
                         plots_irrigated_pct=v26[0], not_irrigated_not_needed_pct=v26[1],
                         not_irrigated_cannot_afford_pct=v26[2], not_irrigated_no_water_pct=v26[3]))
    pd.DataFrame(recs).to_csv(OUT / "ahs_water_province.csv", index=False)

    t38 = openpyxl.load_workbook(RAW / "AHS_2024_Tables.xlsx", data_only=True)["Table 38"]
    soil = []
    for row in t38.iter_rows(min_row=4):
        if row[1].value is None or str(row[1].value).startswith("Source"):
            continue
        for p, c in zip(PROVINCES, row[2:8]):
            soil.append(dict(threat=norm(row[1].value), province=p, pct_households=num(c.value)))
    pd.DataFrame(soil).to_csv(OUT / "ahs_soil_province.csv", index=False)
    print("irrigation rows", len(irr), "| missing farmer %:", int(irr.farmers_irrigating_pct.isna().sum()),
          "| soil rows", len(soil))


if __name__ == "__main__":
    build()
